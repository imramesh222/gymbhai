import datetime as dt
import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.calendar import to_bs
from app.core.time import today_in_nepal
from app.services import checkins
from tests.conftest import (
    GymFixture,
    make_gym,
    make_member,
    priced_plan,
    sell,
    sign_in,
)

TODAY = today_in_nepal()


def check_in(db: Session, fixture: GymFixture, member) -> None:
    checkins.check_in(
        db, gym=fixture.gym, member=member, branch_id=fixture.branch.id, method="manual"
    )
    db.commit()


def test_today(client: TestClient, db: Session, gym_a: GymFixture) -> None:
    paid = make_member(db, gym_a, name="Paid", phone="9800000001")
    sell(db, gym_a, paid, paid=1500_00)
    owing = make_member(db, gym_a, name="Owing", phone="9800000002")
    sell(db, gym_a, owing, paid=500_00, start_date=TODAY - dt.timedelta(days=25),
         end_date=TODAY + dt.timedelta(days=3))  # fmt: skip
    lapsed = make_member(db, gym_a, name="Lapsed", phone="9800000003")
    for m in (paid, lapsed):
        check_in(db, gym_a, m)
    sign_in(db, client, gym_a.owner)

    today = client.get("/api/v1/dashboard/today").json()
    assert today["check_ins"] == 1
    assert [p["name"] for p in today["inside_now"]] == ["Paid"]
    assert [p["name"] for p in today["turned_away"]] == ["Lapsed"]
    assert today["expiring_this_week"] == 1
    assert today["members_with_dues"] == 1 and today["dues_total"] == 1000_00
    assert today["collected"] == 2000_00
    assert today["active_members"] == 2


def test_money_is_hidden_without_the_permission(
    client: TestClient, db: Session, gym_a: GymFixture, make_staff
) -> None:
    sign_in(db, client, make_staff(gym_a.gym, permissions=["members.view"]))
    today = client.get("/api/v1/dashboard/today").json()
    assert today["collected"] is None and today["pending_requests"] is None


def test_monthly_report_in_bs_months(
    client: TestClient, db: Session, gym_a: GymFixture
) -> None:
    # gym_a shows dates in AD and BS, so months are BS months.
    first = make_member(db, gym_a, name="First", phone="9800000011")
    sell(db, gym_a, first, paid=1500_00)
    renewer = make_member(db, gym_a, name="Renewer", phone="9800000012")
    old, _ = sell(db, gym_a, renewer, paid=1500_00)
    sell(
        db, gym_a, renewer, paid=1500_00, start_date=old.end_date + dt.timedelta(days=1)
    )
    sign_in(db, client, gym_a.owner)

    year, month, _ = to_bs(TODAY)
    report = client.get("/api/v1/reports/monthly").json()
    assert report["month"] == f"{year}-{month:02d}" and report["calendar"] == "bs"
    assert report["income"] == 4500_00
    assert report["income_by_method"] == {"cash": 4500_00}
    assert report["new_members"] == 2
    assert report["new_memberships"] == 2 and report["renewals"] == 1
    assert report["by_plan"] == {"1 month": 3}
    assert report["active_members"] == 2

    earlier = client.get(
        "/api/v1/reports/monthly", params={"month": f"{year - 1}-01"}
    ).json()
    assert earlier["income"] == 0
    assert (
        client.get("/api/v1/reports/monthly", params={"month": "2083-13"}).status_code
        == 422
    )


def test_ad_gyms_get_ad_months(
    client: TestClient, db: Session, gym_a: GymFixture
) -> None:
    sign_in(db, client, gym_a.owner)
    client.patch("/api/v1/gym", json={"settings": {"date_display": "ad"}})
    report = client.get("/api/v1/reports/monthly").json()
    assert report["month"] == f"{TODAY:%Y-%m}" and report["calendar"] == "ad"
    assert report["first_day"] == f"{TODAY:%Y-%m}-01"


def workbook(response) -> list[list]:
    sheet = load_workbook(io.BytesIO(response.content)).active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def test_exports(client: TestClient, db: Session, gym_a: GymFixture) -> None:
    member = make_member(db, gym_a)
    sell(db, gym_a, member, paid=1000_00)
    check_in(db, gym_a, member)
    sign_in(db, client, gym_a.owner)

    members = client.get("/api/v1/export/members.xlsx")
    assert members.headers["content-type"].startswith("application/vnd.openxmlformats")
    rows = workbook(members)
    header = rows[0]
    sita = dict(zip(header, rows[1], strict=True))
    assert sita["Name"] == "Sita Rai" and sita["Owes (Rs)"] == 500
    assert "Valid until (BS)" in header

    payments = workbook(client.get("/api/v1/export/payments.xlsx"))
    paid = dict(zip(payments[0], payments[1], strict=True))
    assert paid["Amount (Rs)"] == 1000  # rupees, not paisa
    assert paid["For"] == "Membership"
    assert len(workbook(client.get("/api/v1/export/memberships.xlsx"))) == 2
    assert len(workbook(client.get("/api/v1/export/check-ins.xlsx"))) == 2
    assert client.get("/api/v1/export/secrets.xlsx").status_code == 422


def test_payments_export_needs_money_permission(
    client: TestClient, db: Session, gym_a: GymFixture, make_staff
) -> None:
    sign_in(db, client, make_staff(gym_a.gym, permissions=["members.view"]))
    assert client.get("/api/v1/export/members.xlsx").status_code == 200
    assert client.get("/api/v1/export/payments.xlsx").status_code == 403


def test_today_lists_what_to_do(
    client: TestClient, db: Session, gym_a: GymFixture
) -> None:
    ends = make_member(db, gym_a, name="Ends Today", phone="9800000011")
    sell(db, gym_a, ends, paid=1500_00, start_date=TODAY - dt.timedelta(days=29),
         end_date=TODAY)  # fmt: skip
    small = make_member(db, gym_a, name="Owes A Little", phone="9800000012")
    sell(db, gym_a, small, paid=1400_00)
    big = make_member(db, gym_a, name="Owes A Lot", phone="9800000013")
    sell(db, gym_a, big, paid=0)
    sign_in(db, client, gym_a.owner)

    today = client.get("/api/v1/dashboard/today").json()
    assert today["ending_today"] == [
        {"member_id": str(ends.id), "name": "Ends Today", "phone": "9800000011"}
    ]
    assert [(m["name"], m["dues"]) for m in today["top_owing"]] == [
        ("Owes A Lot", 1500_00),
        ("Owes A Little", 100_00),
    ]


def test_check_ins_by_hour_and_since_yesterday(
    client: TestClient, db: Session, gym_a: GymFixture
) -> None:
    from sqlalchemy import update

    from app.core.time import utcnow
    from app.models.member_app import CheckIn

    early = make_member(db, gym_a, name="Early", phone="9800000021")
    late = make_member(db, gym_a, name="Late", phone="9800000022")
    for m in (early, late):
        sell(db, gym_a, m, paid=1500_00)
        check_in(db, gym_a, m)
    # Yesterday: one before this time of day, one after it.
    now = utcnow()
    db.execute(
        update(CheckIn)
        .where(CheckIn.member_id == early.id)
        .values(at=now - dt.timedelta(days=1, minutes=5))
    )
    db.execute(
        update(CheckIn)
        .where(CheckIn.member_id == late.id)
        .values(at=now - dt.timedelta(days=1) + dt.timedelta(minutes=5))
    )
    today_member = make_member(db, gym_a, name="Today", phone="9800000023")
    sell(db, gym_a, today_member, paid=1500_00)
    check_in(db, gym_a, today_member)
    sign_in(db, client, gym_a.owner)

    today = client.get("/api/v1/dashboard/today").json()
    assert len(today["check_ins_by_hour"]) == 24
    assert sum(today["check_ins_by_hour"]) == today["check_ins"] == 1
    assert today["check_ins_yesterday"] == 1
    assert today["collected_yesterday"] == 0


def test_the_setup_checklist(
    client: TestClient, db: Session, gym_a: GymFixture, make_staff
) -> None:
    sign_in(db, client, gym_a.owner)
    setup = client.get("/api/v1/dashboard/today").json()["setup"]
    assert setup == {
        "plans_priced": False,
        "payment_account": False,
        "logo": False,
        "staff_invited": False,
        "poster_printed": False,
        "first_member": False,
    }

    make_member(db, gym_a)
    client.patch("/api/v1/gym", json={"settings": {"poster_printed": True}})
    setup = client.get("/api/v1/dashboard/today").json()["setup"]
    assert setup["first_member"] and setup["poster_printed"]

    priced_plan(db, gym_a.gym)
    assert client.get("/api/v1/dashboard/today").json()["setup"]["plans_priced"]

    client.patch("/api/v1/gym", json={"settings": {"setup_dismissed": True}})
    assert client.get("/api/v1/dashboard/today").json()["setup"] is None


def test_plans_arent_set_up_until_months_are_counted_somehow(
    client: TestClient, db: Session
) -> None:
    fixture = make_gym(db, slug="new-gym", phone="9841000009", plan_months=None)
    priced_plan(db, fixture.gym)
    sign_in(db, client, fixture.owner)
    assert not client.get("/api/v1/dashboard/today").json()["setup"]["plans_priced"]
    client.patch("/api/v1/gym", json={"settings": {"plan_months": "ad"}})
    assert client.get("/api/v1/dashboard/today").json()["setup"]["plans_priced"]


def test_the_setup_checklist_is_for_whoever_sets_the_gym_up(
    client: TestClient, db: Session, gym_a: GymFixture, make_staff
) -> None:
    sign_in(db, client, make_staff(gym_a.gym, permissions=["members.view"]))
    assert client.get("/api/v1/dashboard/today").json()["setup"] is None


def test_twelve_months_for_the_charts(
    client: TestClient, db: Session, gym_a: GymFixture
) -> None:
    from app.services import reports

    member = make_member(db, gym_a)
    sell(db, gym_a, member, paid=1500_00)
    sign_in(db, client, gym_a.owner)
    body = client.get("/api/v1/reports/trend").json()
    assert body["calendar"] == "bs"
    months = body["months"]
    assert len(months) == 12
    assert [m["month"] for m in months] == reports.months_back(
        reports.current_month("bs"), 12
    )
    this = months[-1]
    assert this["started"] and this["income"] == 1500_00 and this["new_members"] == 1
    # The gym joined this month: the year before is a gap, not zeros to plot.
    assert not months[0]["started"]


def test_months_back_crosses_the_new_year() -> None:
    from app.services import reports

    assert reports.months_back("2083-02", 3) == ["2082-12", "2083-01", "2083-02"]


def test_attendance_by_weekday_and_hour(
    client: TestClient, db: Session, gym_a: GymFixture
) -> None:
    from app.core.time import utcnow

    member = make_member(db, gym_a)
    sell(db, gym_a, member, paid=1500_00)
    check_in(db, gym_a, member)
    sign_in(db, client, gym_a.owner)
    summary = client.get(
        f"/api/v1/check-ins/summary?date_from={TODAY}&date_to={TODAY}"
    ).json()
    grid = summary["by_weekday_hour"]
    assert len(grid) == 7 and all(len(row) == 24 for row in grid)
    from app.core.time import NEPAL

    now = utcnow().astimezone(NEPAL)
    assert grid[(now.weekday() + 1) % 7][now.hour] == 1
    assert sum(map(sum, grid)) == summary["let_in"] == 1


def test_the_trend_agrees_with_each_months_report(db: Session) -> None:
    """Reports' charts read a year in one go (reports.trend); every figure must
    be what the month's own report says, for the owner and for staff who see
    one branch. On the demo gym: a year of members, money, visits, lapses."""
    from dataclasses import astuple, replace

    from sqlalchemy import select

    from app.api.tenancy import StaffContext
    from app.models.gym import Branch, Gym
    from app.models.staff import StaffUser
    from app.services import reports
    from scripts import seed_demo
    from tests.conftest import context_for

    seed_demo.build(db, members=30, say=lambda _: None)
    gym = db.scalar(select(Gym).where(Gym.slug == "demo"))
    owner = db.scalar(
        select(StaffUser).where(StaffUser.gym_id == gym.id, StaffUser.is_owner)
    )
    everyone = context_for(db, owner)
    branch = db.scalar(select(Branch).where(Branch.gym_id == gym.id).limit(1))
    one_branch = StaffContext(
        staff=owner,
        gym_id=gym.id,
        permissions=everyone.permissions,
        branch_ids=frozenset({branch.id}),
    )
    calendar = "ad" if gym.config.date_display == "ad" else "bs"
    months = reports.months_back(reports.current_month(calendar), 14)
    for ctx in (everyone, one_branch):
        batched = reports.trend(db, ctx, months)
        for got, month in zip(batched, months, strict=True):
            expected = reports.monthly(db, ctx, month)
            # The charts don't show the breakdowns; trend leaves them empty.
            expected = replace(
                expected, income_by_method={}, expenses_by_category={}, by_plan={}
            )
            assert astuple(got) == astuple(expected), month
    assert any(m.renewals for m in batched) and any(m.lapsed for m in batched)
