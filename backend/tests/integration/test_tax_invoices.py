"""Tax invoices for VAT-registered gyms (PLAN.md §20.4): one with each
payment, credit notes for voids and refunds, copies counted, the sales book,
and each bill to IRD's CBMS."""

import io
from urllib.parse import urlparse

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.time import today_in_nepal
from app.models.job import Job
from app.models.tax import TaxInvoice
from app.services import tax
from tests.conftest import GymFixture, make_member, sell, sign_in

PAN = "301234567"


def text_of(pdf: bytes) -> str:
    return "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf)).pages)


def path_of(url: str) -> str:
    parts = urlparse(url)
    return f"{parts.path}?{parts.query}"


@pytest.fixture
def vat(client: TestClient, db: Session, gym_a: GymFixture) -> GymFixture:
    sign_in(db, client, gym_a.owner)
    on = client.put(
        "/api/v1/gym/tax",
        json={
            "vat_registered": True,
            "pan": PAN,
            "starts_on": today_in_nepal().isoformat(),
        },
    )
    assert on.status_code == 200, on.text
    return gym_a


def receipt(client: TestClient, payment_id) -> dict:
    return client.get(f"/api/v1/payments/{payment_id}/receipt").json()


def test_off_on_a_server_until_certified(
    client: TestClient, db: Session, gym_a: GymFixture, monkeypatch
) -> None:
    sign_in(db, client, gym_a.owner)
    monkeypatch.setattr(settings, "environment", "production")
    assert client.get("/api/v1/gym/tax").json()["available"] is False
    refused = client.put(
        "/api/v1/gym/tax",
        json={"vat_registered": True, "pan": PAN, "starts_on": "2026-10-20"},
    )
    assert refused.status_code == 409 and refused.json()["code"] == "tax_unavailable"


def test_a_payment_gets_its_invoice(
    client: TestClient, db: Session, vat: GymFixture
) -> None:
    member = make_member(db, vat, name="Sita Rai")
    _, payment = sell(db, vat, member, paid=1500_00)
    [invoice] = receipt(client, payment.id)["tax_documents"]
    assert invoice["kind"] == "invoice"
    assert invoice["number"] == 1
    assert invoice["fiscal_year"] == tax.fiscal_year(today_in_nepal())
    assert invoice["seller"]["pan"] == PAN
    assert invoice["buyer_name"] == "Sita Rai"
    # Prices include VAT: Rs 1,500 is Rs 1,327.43 and Rs 172.57 VAT.
    assert (invoice["taxable"], invoice["vat"], invoice["total"]) == (
        1327_43,
        172_57,
        1500_00,
    )
    assert sum(line["amount"] for line in invoice["lines"]) == 1327_43
    assert "membership" in invoice["lines"][0]["description"]
    assert invoice["total_in_words"] == "One thousand five hundred rupees only"

    _, second = sell(db, vat, make_member(db, vat, phone="9800000002"), paid=1000_00)
    assert receipt(client, second.id)["tax_documents"][0]["number"] == 2


def test_none_before_the_day_it_starts(
    client: TestClient, db: Session, gym_a: GymFixture
) -> None:
    sign_in(db, client, gym_a.owner)
    _, before = sell(db, gym_a, make_member(db, gym_a), paid=1500_00)
    client.put(
        "/api/v1/gym/tax",
        json={"vat_registered": True, "pan": PAN, "starts_on": "2099-01-01"},
    )
    _, still_before = sell(
        db, gym_a, make_member(db, gym_a, phone="9800000003"), paid=1500_00
    )
    assert receipt(client, before.id)["tax_documents"] == []
    assert receipt(client, still_before.id)["tax_documents"] == []


def test_a_void_issues_a_credit_note_and_nothing_is_edited(
    client: TestClient, db: Session, vat: GymFixture
) -> None:
    _, payment = sell(db, vat, make_member(db, vat), paid=1500_00)
    edited = client.patch(f"/api/v1/payments/{payment.id}", json={"amount": 1200_00})
    assert edited.status_code == 409 and edited.json()["code"] == "invoiced"
    voided = client.post(
        f"/api/v1/payments/{payment.id}/void", json={"reason": "Typed twice"}
    )
    assert voided.status_code == 200, voided.text
    invoice, note = receipt(client, payment.id)["tax_documents"]
    assert (note["kind"], note["number"], note["ref_number"]) == ("credit_note", 1, 1)
    assert (note["total"], note["reason"]) == (1500_00, "Typed twice")
    # The invoice itself is as it was.
    assert invoice["total"] == 1500_00


def test_a_refund_issues_a_credit_note_against_its_invoice(
    client: TestClient, db: Session, vat: GymFixture
) -> None:
    membership, payment = sell(db, vat, make_member(db, vat), paid=1500_00)
    refunded = client.post(
        "/api/v1/payments/refunds",
        json={"membership_id": str(membership.id), "amount": 500_00, "method": "cash"},
    )
    assert refunded.status_code == 201, refunded.text
    [note] = receipt(client, refunded.json()["id"])["tax_documents"]
    assert (note["kind"], note["total"], note["ref_number"]) == (
        "credit_note",
        500_00,
        receipt(client, payment.id)["tax_documents"][0]["number"],
    )
    # Voiding the refund puts the money back on the books: a new invoice.
    client.post(
        f"/api/v1/payments/{refunded.json()['id']}/void",
        json={"reason": "Not refunded"},
    )
    kinds = [d["kind"] for d in receipt(client, refunded.json()["id"])["tax_documents"]]
    assert kinds == ["credit_note", "invoice"]


def test_a_shop_sale_lists_what_was_bought(
    client: TestClient, db: Session, vat: GymFixture
) -> None:
    whey = client.post(
        "/api/v1/products",
        json={"name": "Whey 1 kg", "category": "supplement", "price": 4500_00},
    ).json()
    water = client.post(
        "/api/v1/products", json={"name": "Water", "category": "drink", "price": 30_00}
    ).json()
    sold = client.post(
        "/api/v1/sales",
        json={
            "lines": [
                {"product_id": whey["id"], "quantity": 2},
                {"product_id": water["id"], "quantity": 3},
            ],
            "discount": 90_00,
            "payment": {"method": "cash"},
        },
    )
    assert sold.status_code == 201, sold.text
    [invoice] = receipt(client, sold.json()["payment_id"])["tax_documents"]
    assert invoice["buyer_name"] == "Walk-in"
    assert [line["description"] for line in invoice["lines"]] == [
        "Water",
        "Whey 1 kg",
        "Discount",
    ]
    assert invoice["total"] == 2 * 4500_00 + 3 * 30_00 - 90_00
    assert sum(line["amount"] for line in invoice["lines"]) == invoice["taxable"]


def test_the_second_print_is_a_copy(
    client: TestClient, db: Session, vat: GymFixture
) -> None:
    _, payment = sell(db, vat, make_member(db, vat), paid=1500_00)
    first = client.post(f"/api/v1/payments/{payment.id}/receipt/printed").json()
    assert first["tax_documents"][0]["prints"] == 1
    link = client.post(f"/api/v1/payments/{payment.id}/receipt-link").json()
    client.headers.pop("Authorization", None)
    pdf = client.get(path_of(link["url"]))
    text = text_of(pdf.content)
    for expected in ("TAX INVOICE", f"PAN: {PAN}", "Copy of original (2)", "VAT 13%",
                     "Rs 172.57", "One thousand five hundred rupees only"):  # fmt: skip
        assert expected in text, expected
    assert "not a VAT invoice" not in text


def test_the_sales_book(client: TestClient, db: Session, vat: GymFixture) -> None:
    _, kept = sell(db, vat, make_member(db, vat), paid=1500_00)
    _, voided = sell(db, vat, make_member(db, vat, phone="9800000004"), paid=1000_00)
    client.post(f"/api/v1/payments/{voided.id}/void", json={"reason": "Wrong member"})
    book = client.get("/api/v1/tax-invoices").json()
    assert [i["number"] for i in book["invoices"]] == [1, 2]
    assert [r["ref_number"] for r in book["returns"]] == [2]
    assert book["sales_total"] == 2500_00 and book["returns_total"] == 1000_00
    assert book["sales_vat"] + book["sales_taxable"] == book["sales_total"]

    sheet = client.get("/api/v1/export/sales-book.xlsx")
    assert sheet.status_code == 200
    rows = list(
        load_workbook(io.BytesIO(sheet.content)).active.iter_rows(values_only=True)
    )
    assert rows[0][:4] == ("Fiscal year", "Date (BS)", "Date (AD)", "Invoice no.")
    assert [row[3] for row in rows[1:]] == [1, 2]
    returns = client.get("/api/v1/export/sales-returns.xlsx")
    rows = list(
        load_workbook(io.BytesIO(returns.content)).active.iter_rows(values_only=True)
    )
    assert rows[1][4:6] == (2, "Wrong member")


def test_each_bill_goes_to_cbms(
    client: TestClient, db: Session, vat: GymFixture, monkeypatch
) -> None:
    set_up = client.put(
        "/api/v1/gym/tax/cbms", json={"username": "Test_CBMS", "password": "test@321"}
    )
    assert set_up.status_code == 200 and set_up.json()["cbms_enabled"] is True
    _, payment = sell(db, vat, make_member(db, vat), paid=1500_00)
    invoice = db.scalars(select(TaxInvoice)).one()
    assert invoice.cbms_status == "waiting"
    job = db.scalars(select(Job).where(Job.kind == tax.JOB_CBMS)).one()

    sent: list[tuple[str, dict]] = []
    monkeypatch.setattr(
        tax, "_post", lambda path, body: sent.append((path, body)) or 200
    )
    tax.handle_cbms(db, job.payload)
    db.commit()
    db.refresh(invoice)
    assert invoice.cbms_status == "sent"
    [(path, body)] = sent
    first, second = invoice.fiscal_year.split("/")
    assert path == "/api/bill"
    assert body["fiscal_year"] == f"{first}.0{second}"
    assert body["seller_pan"] == PAN and body["password"] == "test@321"
    assert body["invoice_number"] == "1" and body["total_sales"] == 1500.0
    assert body["taxable_sales_vat"] + body["vat"] == pytest.approx(1500.0)
    assert len(body["invoice_date"]) == 10 and body["invoice_date"][4] == "."

    # The credentials refused: the bill isn't sent again, and Settings says why.
    client.post(f"/api/v1/payments/{payment.id}/void", json={"reason": "Typed twice"})
    note = db.scalars(select(TaxInvoice).where(TaxInvoice.kind == "credit_note")).one()
    monkeypatch.setattr(tax, "_post", lambda path, body: 100)
    tax.handle_cbms(db, {"id": str(note.id)})
    db.commit()
    db.refresh(note)
    assert note.cbms_status == "refused"
    settings_now = client.get("/api/v1/gym/tax").json()
    assert settings_now["cbms_error"] == "CBMS refused the username and password."
    assert settings_now["cbms_counts"] == {"sent": 1, "refused": 1}


def test_the_pan_stays_once_invoices_are_issued(
    client: TestClient, db: Session, vat: GymFixture
) -> None:
    sell(db, vat, make_member(db, vat), paid=1500_00)
    changed = client.put(
        "/api/v1/gym/tax",
        json={"vat_registered": True, "pan": "309999999", "starts_on": "2026-01-01"},
    )
    assert changed.status_code == 409 and changed.json()["code"] == "tax_pan_fixed"
    bad = client.put(
        "/api/v1/gym/tax",
        json={"vat_registered": True, "pan": "12345", "starts_on": None},
    )
    assert bad.status_code == 422
