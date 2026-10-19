"""Expenses: money the gym spent (PLAN.md §18.1)."""

import datetime as dt
import io
from collections.abc import Generator

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.calendar import to_bs
from app.core.time import today_in_nepal
from app.models.activity import ActivityLog
from app.models.expense import Expense
from app.models.gym import Branch
from app.services import live, reports
from tests.conftest import GymFixture, lapse, make_member, sell, sign_in

TODAY = today_in_nepal()
URL = "/api/v1/expenses"


def record(client: TestClient, **body) -> dict:
    response = client.post(URL, json={"category": "rent", "amount": 25_000_00, **body})
    assert response.status_code == 201, response.text
    return response.json()


def count(db: Session) -> int:
    return db.scalar(select(func.count()).select_from(Expense)) or 0


def workbook(response) -> list[list]:
    sheet = load_workbook(io.BytesIO(response.content)).active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def test_recording_an_expense(client: TestClient, db: Session, gym_a: GymFixture):
    sign_in(db, client, gym_a.owner)
    expense = record(client, method="bank", paid_to="  Hari Shrestha (landlord) ")
    assert expense["spent_on"] == TODAY.isoformat()
    assert expense["paid_to"] == "Hari Shrestha (landlord)"
    # No branch given by someone who works at every branch: the whole gym's.
    assert expense["branch_id"] is None
    assert expense["created_by_name"] == gym_a.owner.name

    entry = db.scalars(
        select(ActivityLog).where(ActivityLog.action == "expense.created")
    ).one()
    assert entry.actor_id == gym_a.owner.id
    assert entry.changes["amount"] == {"before": None, "after": 25_000_00}


def test_what_is_refused(client: TestClient, db: Session, gym_a: GymFixture):
    sign_in(db, client, gym_a.owner)
    tomorrow = (TODAY + dt.timedelta(days=1)).isoformat()
    for body, code in [
        ({"category": "rent", "amount": 0}, "validation_error"),
        ({"category": "bribes", "amount": 100}, "validation_error"),
        ({"category": "rent", "amount": 100, "spent_on": tomorrow}, "future_date"),
        # The gym is the token's, never the body's.
        (
            {"category": "rent", "amount": 100, "gym_id": str(gym_a.gym.id)},
            "validation_error",
        ),
    ]:
        response = client.post(URL, json=body)
        assert response.status_code == 422, body
        assert response.json()["code"] == code, body
    assert count(db) == 0


def test_a_double_tap_records_it_once(
    client: TestClient, db: Session, gym_a: GymFixture
):
    sign_in(db, client, gym_a.owner)
    body = {"category": "electricity", "amount": 3_200_00}
    headers = {"Idempotency-Key": "expense-electricity-1"}
    first = client.post(URL, json=body, headers=headers)
    again = client.post(URL, json=body, headers=headers)
    assert first.status_code == again.status_code == 201
    assert again.headers["Idempotent-Replayed"] == "true"
    assert again.json() == first.json()
    assert count(db) == 1


def test_a_month_of_expenses_with_its_totals(
    client: TestClient, db: Session, gym_a: GymFixture
):
    """BS months for a gym that shows BS; voided ones are listed, not counted."""
    sign_in(db, client, gym_a.owner)
    record(client, category="rent", amount=25_000_00)
    record(client, category="salaries", amount=18_000_00, paid_to="Ram (trainer)")
    record(client, category="salaries", amount=15_000_00, paid_to="Gita (trainer)")
    mistake = record(client, category="water", amount=500_00)
    voided = client.post(f"{URL}/{mistake['id']}/void", json={"reason": "Twice"})
    assert voided.status_code == 200, voided.text

    year, month, _ = to_bs(TODAY)
    this_month = f"{year}-{month:02d}"
    first_day, _ = reports.month_range(this_month, "bs")
    last_month = record(
        client,
        category="internet",
        amount=1_200_00,
        spent_on=(first_day - dt.timedelta(days=1)).isoformat(),
    )

    listing = client.get(URL).json()
    assert (listing["month"], listing["calendar"]) == (this_month, "bs")
    assert listing["first_day"] == first_day.isoformat()
    assert len(listing["items"]) == 4
    assert listing["sum_amount"] == 58_000_00
    assert listing["by_category"] == {"rent": 25_000_00, "salaries": 33_000_00}
    gone = next(i for i in listing["items"] if i["id"] == mistake["id"])
    assert gone["void_reason"] == "Twice"

    salaries = client.get(URL, params={"category": "salaries"}).json()
    assert sorted(i["paid_to"] for i in salaries["items"]) == [
        "Gita (trainer)",
        "Ram (trainer)",
    ]
    assert salaries["sum_amount"] == 33_000_00
    assert client.get(URL, params={"category": "bribes"}).status_code == 422

    before = reports.months_back(this_month, 2)[0]
    earlier = client.get(URL, params={"month": before}).json()
    assert [i["id"] for i in earlier["items"]] == [last_month["id"]]
    assert client.get(URL, params={"month": "2083-13"}).status_code == 422


def test_ad_gyms_list_ad_months(client: TestClient, db: Session, gym_a: GymFixture):
    sign_in(db, client, gym_a.owner)
    client.patch("/api/v1/gym", json={"settings": {"date_display": "ad"}})
    record(client)
    listing = client.get(URL).json()
    assert (listing["month"], listing["calendar"]) == (f"{TODAY:%Y-%m}", "ad")
    assert listing["first_day"] == f"{TODAY:%Y-%m}-01"
    assert listing["sum_amount"] == 25_000_00


def test_edits_are_logged_with_before_and_after(
    client: TestClient, db: Session, gym_a: GymFixture
):
    sign_in(db, client, gym_a.owner)
    expense = record(client, category="repairs", amount=2_000_00)
    response = client.patch(
        f"{URL}/{expense['id']}",
        json={
            "amount": 2_500_00,
            "note": "Treadmill belt",
            # Null for a field that always has a value leaves it alone.
            "category": None,
            "reason": "Typed the wrong amount",
        },
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert (body["category"], body["amount"], body["note"]) == (
        "repairs",
        2_500_00,
        "Treadmill belt",
    )
    updates = select(ActivityLog).where(ActivityLog.action == "expense.updated")
    entry = db.scalars(updates).one()
    assert entry.changes == {
        "amount": {"before": 2_000_00, "after": 2_500_00},
        "note": {"before": None, "after": "Treadmill belt"},
    }
    assert entry.reason == "Typed the wrong amount"

    # Nothing changed, nothing logged.
    client.patch(f"{URL}/{expense['id']}", json={"amount": 2_500_00})
    assert len(db.scalars(updates).all()) == 1
    tomorrow = (TODAY + dt.timedelta(days=1)).isoformat()
    later = client.patch(f"{URL}/{expense['id']}", json={"spent_on": tomorrow})
    assert later.json()["code"] == "future_date"


def test_voiding(client: TestClient, db: Session, gym_a: GymFixture):
    sign_in(db, client, gym_a.owner)
    url = f"{URL}/{record(client)['id']}"
    assert client.post(f"{url}/void", json={"reason": "no"}).status_code == 422
    voided = client.post(f"{url}/void", json={"reason": "The landlord paid it"})
    assert voided.json()["voided_at"] is not None

    again = client.post(f"{url}/void", json={"reason": "The landlord paid it"})
    assert (again.status_code, again.json()["code"]) == (409, "expense_voided")
    edit = client.patch(url, json={"amount": 1})
    assert (edit.status_code, edit.json()["code"]) == (409, "expense_voided")
    entry = db.scalars(
        select(ActivityLog).where(ActivityLog.action == "expense.voided")
    ).one()
    assert entry.reason == "The landlord paid it"
    assert client.get("/api/v1/reports/monthly").json()["expenses"] == 0
    assert count(db) == 1  # voided, never deleted


def test_staff_at_one_branch(
    client: TestClient, db: Session, gym_a: GymFixture, make_staff
):
    koteshwor = Branch(gym_id=gym_a.gym.id, name="Koteshwor")
    db.add(koteshwor)
    db.commit()
    sign_in(db, client, gym_a.owner)
    record(client, amount=40_000_00)  # the whole gym's
    other = record(client, amount=20_000_00, branch_id=str(koteshwor.id))
    theirs = record(client, amount=10_000_00, branch_id=str(gym_a.branch.id))

    desk = make_staff(
        gym_a.gym,
        permissions=["reports.money", "expenses.record", "expenses.edit"],
        branches=[gym_a.branch],
    )
    sign_in(db, client, desk)
    listing = client.get(URL).json()
    assert [i["id"] for i in listing["items"]] == [theirs["id"]]
    assert listing["sum_amount"] == 10_000_00
    assert client.get("/api/v1/reports/monthly").json()["expenses"] == 10_000_00

    for branch_id in (None, str(koteshwor.id)):
        refused = client.post(
            URL, json={"category": "rent", "amount": 100, "branch_id": branch_id}
        )
        assert (refused.status_code, refused.json()["code"]) == (403, "branch_denied")
    mine = record(client, amount=100, branch_id=str(gym_a.branch.id))
    assert mine["branch_id"] == str(gym_a.branch.id)

    # Another branch's expense is not found, as another gym's would be ...
    assert client.patch(f"{URL}/{other['id']}", json={"note": "x"}).status_code == 404
    # ... and theirs can't be made the whole gym's.
    moved = client.patch(f"{URL}/{theirs['id']}", json={"branch_id": None})
    assert (moved.status_code, moved.json()["code"]) == (403, "branch_denied")


def test_salaries_are_only_for_staff_who_see_money(
    client: TestClient, db: Session, gym_a: GymFixture, make_staff
):
    """Recording is allowed without it; the list and totals are not (§18.1)."""
    sign_in(db, client, make_staff(gym_a.gym, permissions=["expenses.record"]))
    recorded = client.post(URL, json={"category": "cleaning", "amount": 800_00})
    assert recorded.status_code == 201
    listed = client.get(URL)
    assert listed.status_code == 403
    assert listed.json()["permission"] == "reports.money"

    sign_in(db, client, make_staff(gym_a.gym, permissions=["members.view"]))
    assert client.get("/api/v1/export/expenses.xlsx").status_code == 403


def test_reports_show_spending_and_profit(
    client: TestClient, db: Session, gym_a: GymFixture
):
    member = make_member(db, gym_a)
    sell(db, gym_a, member, paid=1500_00)
    sign_in(db, client, gym_a.owner)
    record(client, category="electricity", amount=1_000_00)
    record(client, category="rent", amount=2_000_00)

    report = client.get("/api/v1/reports/monthly").json()
    assert (report["income"], report["expenses"]) == (1500_00, 3_000_00)
    assert report["expenses_by_category"] == {"electricity": 1_000_00, "rent": 2_000_00}
    # A month that spent more than came in.
    assert report["profit"] == -1_500_00

    months = client.get("/api/v1/reports/trend").json()["months"]
    assert (months[-1]["expenses"], months[-1]["profit"]) == (3_000_00, -1_500_00)
    assert months[-2]["expenses"] == 0


def test_the_export(client: TestClient, db: Session, gym_a: GymFixture):
    sign_in(db, client, gym_a.owner)
    record(client, category="salaries", amount=18_000_00, paid_to="Ram (trainer)")
    rows = workbook(client.get("/api/v1/export/expenses.xlsx"))
    assert "Date (BS)" in rows[0]
    row = dict(zip(rows[0], rows[1], strict=True))
    assert row["Date (AD)"] == TODAY.isoformat()
    assert (row["Category"], row["Amount (Rs)"], row["Paid to"]) == (
        "salaries",
        18_000,
        "Ram (trainer)",
    )
    assert (row["Branch"], row["Recorded by"]) == ("Whole gym", gym_a.owner.name)
    assert row["Voided"] is False


def test_read_only_once_the_subscription_lapses(
    client: TestClient, db: Session, gym_a: GymFixture
):
    sign_in(db, client, gym_a.owner)
    lapse(db, gym_a.gym)
    response = client.post(URL, json={"category": "rent", "amount": 100})
    assert (response.status_code, response.json()["code"]) == (
        402,
        "subscription_lapsed",
    )
    assert client.get(URL).status_code == 200


@pytest.fixture
def published(monkeypatch) -> Generator[list[dict]]:
    calls: list[dict] = []
    monkeypatch.setattr(live, "publish", lambda db, **kwargs: calls.append(kwargs))
    yield calls


def test_changes_are_published(
    client: TestClient, db: Session, gym_a: GymFixture, published: list[dict]
):
    sign_in(db, client, gym_a.owner)
    expense = record(client, branch_id=str(gym_a.branch.id))
    assert published == [
        {"gym_id": gym_a.gym.id, "kind": live.EXPENSE, "branch_id": gym_a.branch.id}
    ]

    # Made the whole gym's: the branch it left hears about it too.
    published.clear()
    client.patch(f"{URL}/{expense['id']}", json={"branch_id": None})
    assert [c["branch_id"] for c in published] == [None, gym_a.branch.id]

    published.clear()
    client.post(f"{URL}/{expense['id']}/void", json={"reason": "A mistake"})
    assert [c["kind"] for c in published] == [live.EXPENSE]
