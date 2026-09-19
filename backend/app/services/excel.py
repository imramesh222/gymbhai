"""Excel export and register import (PLAN.md §1 "it's their data", §7).

Export: members, memberships, payments and check-ins, whenever the gym wants
— including when its subscription has lapsed.

Import: an existing register (.xlsx or .csv) becomes members with their
current expiry dates. Columns are matched by name and can be re-mapped before
anything is saved. Dates may be AD or BS: a year of 2050 or more is BS.
"""

import csv
import datetime as dt
import io
import re
from collections.abc import Iterable, Sequence
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.core.calendar import from_bs
from app.core.errors import AppError
from app.core.phone import normalize_phone

MAX_ROWS = 5000
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Field -> words that suggest a column is that field.
FIELDS: dict[str, tuple[str, ...]] = {
    "name": ("name", "member", "naam"),
    "phone": ("phone", "mobile", "contact", "cell", "number"),
    "email": ("email", "e-mail", "mail"),
    "gender": ("gender", "sex"),
    "member_code": ("code", "id", "card"),
    "joined_on": ("joined", "join", "admission", "registered"),
    "plan_name": ("plan", "package", "membership type", "type"),
    "start_date": ("start", "from", "renewed", "paid on"),
    "end_date": ("end", "expiry", "expires", "valid", "till", "until", "to"),
    "notes": ("note", "remark", "comment"),
}


# --- export ------------------------------------------------------------------


def workbook(
    title: str, headers: Sequence[str], rows: Iterable[Sequence[Any]]
) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = title[:31]
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    widths = [len(h) for h in headers]
    for row in rows:
        sheet.append(list(row))
        for i, value in enumerate(row):
            widths[i] = min(max(widths[i], len(str(value or ""))), 50)
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width + 2
    sheet.freeze_panes = "A2"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def rupees(paisa: int | None) -> float | None:
    return None if paisa is None else paisa / 100


# --- import ------------------------------------------------------------------


def _cell(value: Any) -> Any:
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return "" if value is None else str(value).strip()


def read_table(filename: str, data: bytes) -> tuple[list[str], list[list[Any]]]:
    """Headers and rows from an .xlsx or .csv file."""
    name = filename.lower()
    try:
        if name.endswith(".csv"):
            text = data.decode("utf-8-sig", errors="replace")
            table = list(csv.reader(io.StringIO(text)))
        elif name.endswith((".xlsx", ".xlsm")):
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            table = [list(r) for r in book.worksheets[0].iter_rows(values_only=True)]
        else:
            raise AppError(415, "unsupported_file", "Upload an .xlsx or .csv file.")
    except AppError:
        raise
    except Exception as exc:
        raise AppError(422, "unreadable_file", "That file couldn't be read.") from exc
    table = [[_cell(v) for v in row] for row in table if any(_cell(v) for v in row)]
    if not table:
        raise AppError(422, "empty_file", "The file has no rows.")
    headers = [str(h) or f"Column {i + 1}" for i, h in enumerate(table[0])]
    rows = [r + [""] * (len(headers) - len(r)) for r in table[1 : MAX_ROWS + 1]]
    return headers, [r[: len(headers)] for r in rows]


def suggest_mapping(headers: Sequence[str]) -> dict[str, int | None]:
    """Field -> column index, from the column names."""
    mapping: dict[str, int | None] = {}
    taken: set[int] = set()
    for field_name, words in FIELDS.items():
        mapping[field_name] = None
        for i, header in enumerate(headers):
            h = header.lower()
            if i not in taken and any(
                re.search(rf"\b{re.escape(w)}", h) for w in words
            ):
                mapping[field_name] = i
                taken.add(i)
                break
    return mapping


_DATE_PATTERNS = (
    (re.compile(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})$"), ("y", "m", "d")),
    (re.compile(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})$"), ("d", "m", "y")),
)


def parse_date(value: str) -> dt.date | None:
    """AD or BS, as ISO or day-first. A year of 2050 or more is BS."""
    value = (value or "").strip()
    if not value:
        return None
    value = value.split("T")[0].split(" ")[0]
    for pattern, order in _DATE_PATTERNS:
        match = pattern.match(value)
        if match:
            parts = dict(zip(order, (int(g) for g in match.groups()), strict=True))
            try:
                if parts["y"] >= 2050:
                    return from_bs(parts["y"], parts["m"], parts["d"])
                return dt.date(parts["y"], parts["m"], parts["d"])
            except (ValueError, KeyError, IndexError) as exc:
                raise ValueError(f"'{value}' isn't a real date") from exc
    raise ValueError(
        f"'{value}' isn't a date we understand (use 2026-09-19 or 2083-06-03)"
    )


def row_to_member(row: Sequence[Any], mapping: dict[str, int | None]) -> dict[str, Any]:
    """One register row, checked. Raises ValueError with a readable reason."""

    def get(field_name: str) -> str:
        index = mapping.get(field_name)
        return str(row[index]).strip() if index is not None and index < len(row) else ""

    name = get("name")
    if not name:
        raise ValueError("no name")
    phone_raw = get("phone")
    if not phone_raw:
        raise ValueError("no phone number")
    try:
        phone = normalize_phone(phone_raw)
    except ValueError as exc:
        raise ValueError(f"phone '{phone_raw}' isn't a Nepali mobile number") from exc
    gender = get("gender").lower()[:1]
    end = parse_date(get("end_date"))
    start = parse_date(get("start_date"))
    if end and start and end < start:
        raise ValueError("the end date is before the start date")
    return {
        "name": name[:120],
        "phone": phone,
        "email": get("email").lower() or None,
        "gender": {"m": "male", "f": "female"}.get(gender),
        "joined_on": parse_date(get("joined_on")),
        "notes": get("notes")[:2000] or None,
        "plan_name": get("plan_name")[:80] or None,
        "start_date": start,
        "end_date": end,
    }
